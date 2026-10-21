"""Los reportes: detalle, resumen por proveedor, resumen por cliente e incidencias, y el Excel que los junta.

Cada reporte devuelve (columnas, filas); los usan el Excel y el bot. `cliente` (RFC receptor) limita el reporte
a las facturas de ese cliente de KMG; sin él se incluyen todas.
"""
from datetime import datetime
from pathlib import Path

from . import config
from .base import SIN_DATOS, leer, ruta_guardada

# solo facturas de ingreso: una nota de crédito (E) o un complemento de pago (P) no es "facturado"
INGRESO = "tipo='I'"


def donde(cliente, *condiciones):
    """WHERE y parámetros con las condiciones dadas y, si hay cliente, el filtro por receptor."""
    condiciones = list(condiciones) + (["receptor_rfc=?"] if cliente else [])
    return ("WHERE " + " AND ".join(condiciones) if condiciones else ""), ((cliente,) if cliente else ())


def detalle(cliente=None):
    """Una fila por factura."""
    w, p = donde(cliente)
    filas = leer("SELECT uuid, emisor_rfc, emisor_nombre, fecha, subtotal, iva, total, receptor_rfc, receptor_nombre, "
                 f"serie, folio, moneda, archivo FROM facturas {w} ORDER BY fecha, uuid", p)
    return (["UUID", "RFC emisor", "Nombre emisor", "Fecha", "Subtotal", "IVA trasladado", "Total",
             "RFC receptor", "Nombre receptor", "Serie", "Folio", "Moneda", "Archivo"],
            [(*f[:12], Path(f[12]).name) for f in filas])


def resumen(cliente=None):
    """Total facturado e IVA por RFC emisor, de mayor a menor."""
    w, p = donde(cliente, INGRESO)
    return (["RFC emisor", "Nombre emisor", "Facturas", "Subtotal", "IVA trasladado", "Total facturado"],
            leer("SELECT emisor_rfc, MAX(emisor_nombre), COUNT(*), ROUND(SUM(subtotal), 2), ROUND(SUM(iva), 2), "
                 f"ROUND(SUM(total), 2) FROM facturas {w} GROUP BY emisor_rfc ORDER BY 6 DESC", p))


def resumen_clientes():
    """Por cada cliente de KMG (RFC receptor): facturas recibidas, proveedores, IVA y total, de mayor a menor."""
    return (["RFC cliente", "Nombre cliente", "Facturas", "Proveedores", "IVA trasladado", "Total facturado"],
            leer("SELECT receptor_rfc, MAX(receptor_nombre), COUNT(*), COUNT(DISTINCT emisor_rfc), "
                 f"ROUND(SUM(iva), 2), ROUND(SUM(total), 2) FROM facturas WHERE {INGRESO} "
                 "GROUP BY receptor_rfc ORDER BY 6 DESC"))


def incidencias(cliente=None, carpeta=None, sin_dueno=True):
    """Archivos que no se cargaron (ERROR) y facturas cargadas que conviene revisar (AVISO).

    Con cliente: las de sus facturas y, si `sin_dueno`, también las de archivos que no se pudieron leer, que no se
    sabe de quién son. Con carpeta: solo las de los archivos de esa carpeta (las de una carga).
    """
    condiciones, p = [], []
    if cliente:
        suyas = "uuid IN (SELECT uuid FROM facturas WHERE receptor_rfc=?)"
        condiciones.append(f"(uuid IS NULL OR {suyas})" if sin_dueno else suyas)
        p.append(cliente)
    if carpeta:
        prefijo = ruta_guardada(carpeta) + "/"
        condiciones.append("substr(archivo, 1, ?) = ?")
        p += [len(prefijo), prefijo]
    filtro = "WHERE " + " AND ".join(condiciones) if condiciones else ""
    filas = leer(f"SELECT archivo, nivel, tipo, detalle FROM incidencias {filtro} ORDER BY nivel DESC, archivo", p)
    return ["Archivo", "Nivel", "Tipo", "Detalle"], [(Path(f[0]).name, *f[1:]) for f in filas]  # ERROR antes que AVISO


def clientes():
    """[(rfc, nombre, facturas)] de los clientes con facturas en la base, el de más facturas primero."""
    return leer("SELECT receptor_rfc, MAX(receptor_nombre), COUNT(*) FROM facturas WHERE receptor_rfc IS NOT NULL "
                "GROUP BY receptor_rfc ORDER BY 3 DESC, 2")


# --- Excel

# Formato corporativo: la paleta de KM&G (carbón y dorado). Cambiar aquí cambia todo el Excel.
CARBON, DORADO, GRIS, FRANJA, LINEA = "1F1F1F", "C9A13B", "6B7280", "F6F4EF", "DDD8CC"
FUENTE = "Calibri"
FILA_ENCABEZADO = 4  # filas 1-2: título y datos del reporte; 3: separación
DINERO = '"$"#,##0.00'
FECHA = "dd/mm/yyyy hh:mm"


def libro(cliente=None):
    """El Excel de revisión (hojas "Detalle" y "Resumen por proveedor"), o None si no hay facturas."""
    from openpyxl import Workbook

    columnas, facturas = detalle(cliente)
    if not facturas:
        return None
    nombre_cliente = f"{facturas[0][8] or cliente} ({cliente})" if cliente else "Todos"
    generado = f"Generado el {datetime.now():%d/%m/%Y %H:%M} · Cliente: {nombre_cliente}"

    wb = Workbook()
    wb.properties.title = "Revisión de CFDI"
    wb.properties.creator = "Revisión de CFDI"
    filas = []
    for f in facturas:
        f = list(f)
        try:
            f[3] = datetime.fromisoformat(f[3])  # la fecha como fecha de Excel
        except ValueError:
            pass  # se deja el texto tal como venía en el XML
        filas.append(f)
    hoja(wb.active, "Detalle", "Revisión de CFDI · Detalle de facturas",
         f"{generado} · {len(filas)} facturas", columnas, filas, dinero=(4, 5, 6), fecha=3, totales=(4, 5, 6))

    columnas, proveedores = resumen(cliente)
    hoja(wb.create_sheet(), "Resumen por proveedor", "Revisión de CFDI · Resumen por proveedor",
         f"{generado} · {len(proveedores)} proveedores · de mayor a menor total facturado",
         columnas, proveedores, dinero=(3, 4, 5), totales=(2, 3, 4, 5))
    return wb


def hoja(ws, nombre, titulo, subtitulo, columnas, filas, dinero=(), fecha=None, totales=()):
    """Llena una hoja con formato corporativo: título, encabezado oscuro con línea dorada, filas alternadas,
    fila de totales, filtros, primera fila fija y lista para imprimir. `totales` son las columnas que se suman."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ws.title = nombre
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = DORADO
    ultima = get_column_letter(len(columnas))
    linea = Side(style="thin", color=LINEA)
    dorada = Side(style="medium", color=DORADO)

    # título y datos del reporte
    ws["A1"] = titulo
    ws["A1"].font = Font(name=FUENTE, size=16, bold=True, color=CARBON)
    ws["A2"] = subtitulo
    ws["A2"].font = Font(name=FUENTE, size=10, color=GRIS)
    ws.row_dimensions[1].height = 26
    for i in range(1, len(columnas) + 1):  # la línea dorada bajo el título, a todo el ancho de la tabla
        ws.cell(2, i).border = Border(bottom=dorada)

    # encabezado de la tabla
    ws.append([])  # fila 3, separación
    ws.append(columnas)
    for c in ws[FILA_ENCABEZADO]:
        c.font = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=CARBON)
        c.border = Border(bottom=dorada)
        c.alignment = Alignment(horizontal="right" if c.column - 1 in dinero + totales else "left",
                                vertical="center", wrap_text=True, indent=1)
    ws.row_dimensions[FILA_ENCABEZADO].height = 22

    # datos, con filas alternadas
    for n, fila in enumerate(filas):
        ws.append(list(fila))
        for c in ws[ws.max_row]:
            i = c.column - 1
            c.font = Font(name=FUENTE, size=10, color=CARBON)
            c.border = Border(bottom=linea)
            if n % 2:
                c.fill = PatternFill("solid", fgColor=FRANJA)
            if i in dinero:
                c.number_format = DINERO
            elif i == fecha:
                c.number_format = FECHA
            elif isinstance(c.value, int):
                c.number_format = "#,##0"
            # margen interior para que las columnas no se peguen; cifras a la derecha, lo demás a la izquierda
            c.alignment = Alignment(horizontal="right" if i in dinero + totales else "left", indent=1)
    primera, fin = FILA_ENCABEZADO + 1, FILA_ENCABEZADO + len(filas)

    # totales: fórmulas, para que sigan cuadrando si alguien edita el Excel
    if totales and filas:
        ws.append(["TOTAL"] + [None] * (len(columnas) - 1))
        for i in totales:
            letra = get_column_letter(i + 1)
            ws.cell(ws.max_row, i + 1, f"=SUM({letra}{primera}:{letra}{fin})")
        for c in ws[ws.max_row]:
            c.font = Font(name=FUENTE, size=10, bold=True, color=CARBON)
            c.border = Border(top=dorada, bottom=dorada)
            c.number_format = DINERO if c.column - 1 in dinero else "#,##0"
            c.alignment = Alignment(horizontal="right" if c.column - 1 in totales else "left", indent=1)

    # anchos según el contenido
    for i, titulo_col in enumerate(columnas, start=1):
        valores = [ws.cell(r, i).value for r in range(primera, fin + 1)]
        largo = max([len(str(titulo_col))] + [17 if isinstance(v, datetime) else
                                             len(f"{v:,.2f}") + 2 if isinstance(v, float) else
                                             len(str(v or "")) for v in valores])
        ws.column_dimensions[get_column_letter(i)].width = min(max(largo + 4, 10), 45)  # +4: texto + margen interior

    ws.freeze_panes = f"A{primera}"
    ws.auto_filter.ref = f"A{FILA_ENCABEZADO}:{ultima}{fin}"

    # impresión: horizontal, al ancho de la hoja, encabezado en cada página y número de página
    ws.page_setup.orientation = "landscape"
    if len(columnas) > 8:  # el detalle es ancho: en oficio la letra sale más grande que en carta
        ws.page_setup.paperSize = ws.PAPERSIZE_LEGAL
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{FILA_ENCABEZADO}:{FILA_ENCABEZADO}"
    ws.print_options.horizontalCentered = True
    ws.oddFooter.left.text = titulo
    ws.oddFooter.right.text = "Página &P de &N"  # &P y &N son códigos de Excel: página actual y total


def excel(destino=None, cliente=None):
    """Guarda el Excel y devuelve su ruta. Lanza ValueError si no hay facturas y PermissionError si está abierto."""
    destino = Path(destino or config.SALIDA / "reporte_cfdi.xlsx")
    wb = libro(cliente)
    if not wb:
        raise ValueError(SIN_DATOS)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino

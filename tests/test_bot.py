"""Pruebas del bot (bot.py y la carpeta asistente/). La IA se imita con ModeloFalso: no se llama a OpenAI."""
import json
import sqlite3
from io import BytesIO
from types import SimpleNamespace

import pytest

from asistente import preguntas
from asistente.pdf import generar_pdf
from revision import config

IVA_MAYOR = "SELECT emisor_nombre, SUM(iva) FROM facturas GROUP BY emisor_rfc ORDER BY 2 DESC LIMIT 1"
FACTURA_01 = "A5AEC797-8306-D03B-F38B-2FFC80A4DF5A"
OTRO_CLIENTE = "OCD200101XY9"  # receptor de los XML de ejemplos/otro_cliente/
FACTURA_OTRO = "0C1E4A7B-2D3F-4A5B-8C6D-7E8F9A0B1C2D"  # ejemplos/otro_cliente/otro_cliente_01.xml
EJEMPLOS = config.RAIZ / "ejemplos" / "otro_cliente"


class ModeloFalso:
    """Imita a la API de OpenAI: hace las llamadas a herramientas que se le den, en orden, y luego contesta."""

    def __init__(self, llamadas, respuesta="Listo."):
        self.pasos = [self.llamada(n, a) for n, a in llamadas] + [SimpleNamespace(content=respuesta, tool_calls=None)]
        self.recibido = []
        self.chat = SimpleNamespace(completions=SimpleNamespace(create=self.create))

    @staticmethod
    def llamada(nombre, args):
        funcion = SimpleNamespace(name=nombre, arguments=json.dumps(args))
        return SimpleNamespace(content=None, tool_calls=[SimpleNamespace(id="1", function=funcion)])

    def create(self, messages, **_):
        self.recibido.append(list(messages))
        return SimpleNamespace(choices=[SimpleNamespace(message=self.pasos.pop(0))])


def test_pdf(salida):
    archivo = generar_pdf(FACTURA_01.lower())
    assert archivo.parent == salida / "pdf" and archivo.read_bytes().startswith(b"%PDF")
    assert generar_pdf("NO-EXISTE") is None


@pytest.mark.parametrize("sql", ["DELETE FROM facturas", "SELECT 1; DROP TABLE facturas", "PRAGMA table_info(facturas)",
                                 "ATTACH 'otra.db' AS x"])
def test_solo_consultas_de_lectura(sql):
    with pytest.raises(ValueError):
        preguntas.validar(sql)


def test_conexion_de_solo_lectura(salida):
    # aunque algo pasara la validación, la conexión no permite escribir
    con = preguntas.conectar()
    with pytest.raises(sqlite3.OperationalError):
        con.execute("DELETE FROM facturas")
    con.close()


def test_pregunta_en_lenguaje_natural(consulta, monkeypatch):
    modelo = ModeloFalso([("consultar", {"sql": "DROP TABLE facturas"}), ("consultar", {"sql": IVA_MAYOR})],
                         "Papelería del Bajío.")
    monkeypatch.setattr(preguntas.openai, "OpenAI", lambda: modelo)
    respuesta, consultas, adjuntos = preguntas.responder("¿Qué proveedor me cobró más IVA?")
    assert respuesta == "Papelería del Bajío."
    assert "error" in consultas[0]  # el DROP se rechazó y el error regresó al modelo
    assert consultas[1]["filas"] == [("PAPELERIA DEL BAJIO DEMO", 14872.54)]
    assert adjuntos == []
    assert consulta("SELECT COUNT(*) FROM facturas") == [(14,)]


def test_adjuntos(salida, monkeypatch):
    modelo = ModeloFalso([("enviar_pdf", {"uuids": [FACTURA_01.lower(), "NO-EXISTE"]}), ("enviar_excel", {})])
    monkeypatch.setattr(preguntas.openai, "OpenAI", lambda: modelo)
    _, _, adjuntos = preguntas.responder("Dame el PDF de la A1000 y el Excel")
    assert [(a["tipo"], a.get("uuid")) for a in adjuntos] == [("pdf", FACTURA_01), ("xlsx", None)]
    resultado_pdf = modelo.recibido[1][-1]["content"]  # lo que el modelo recibió de enviar_pdf
    assert "NO-EXISTE" in resultado_pdf


@pytest.fixture
def cliente(salida, monkeypatch):
    from asistente import web
    monkeypatch.setenv("OPENAI_API_KEY", "prueba")
    web.historial.clear()
    return web.app.test_client()


def test_botones_del_reto_sin_ia(cliente, monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY")  # funcionan aunque no haya llave
    pagina = cliente.get("/").get_data(as_text=True)
    assert "14 facturas" in pagina and "Sin OPENAI_API_KEY" in pagina

    r = cliente.post("/reto/cargar", json={}).get_json()
    assert "18 archivos XML" in r["texto"] and "14 facturas cargadas" in r["texto"]
    assert r["consultas"][0]["titulo"] == "Incidencias de esta carga"
    assert [c["rfc"] for c in r["clientes"]] == ["CDE150101AB1"]
    r = cliente.post("/reto/detalle", json={}).get_json()
    assert r["consultas"][0]["total"] == 14
    r = cliente.post("/reto/resumen", json={}).get_json()
    assert r["consultas"][0]["filas"][0][0] == "PAP180312KL4"
    r = cliente.post("/reto/incidencias", json={}).get_json()
    assert "4 archivos con error" in r["texto"]
    assert cliente.post("/reto/borrar", json={}).status_code == 404

    assert cliente.get("/reporte_cfdi.xlsx").data[:2] == b"PK"  # un .xlsx es un zip
    assert b"Duplicada" in cliente.get("/incidencias.log").data
    assert "OPENAI_API_KEY" in cliente.post("/preguntar", json={"pregunta": "hola"}).get_json()["texto"]


def test_chat_con_ia(cliente, monkeypatch):
    modelos = [ModeloFalso([("consultar", {"sql": IVA_MAYOR})], "Papelería del Bajío."),
               ModeloFalso([("enviar_pdf", {"uuids": [FACTURA_01]})], "Aquí está.")]
    monkeypatch.setattr(preguntas.openai, "OpenAI", lambda: modelos[0] if modelos[0].pasos else modelos[1])

    r = cliente.post("/preguntar", json={"pregunta": "¿Qué proveedor me cobró más IVA?"}).get_json()
    assert r["texto"] == "Papelería del Bajío."
    assert r["consultas"][0]["sql"] == IVA_MAYOR and r["consultas"][0]["filas"] == [["PAPELERIA DEL BAJIO DEMO", 14872.54]]

    r = cliente.post("/preguntar", json={"pregunta": "¿Y su factura más reciente en PDF?"}).get_json()
    assert r["adjuntos"][0]["url"] == f"/pdf/{FACTURA_01}"
    # la segunda pregunta llevó la primera como historial, con los datos que se consultaron
    previos = [m["content"] for m in modelos[1].recibido[0] if m["role"] == "assistant"]
    assert "PAPELERIA DEL BAJIO DEMO" in previos[0]

    assert cliente.get(r["adjuntos"][0]["url"]).data.startswith(b"%PDF")
    assert cliente.get("/pdf/NO-EXISTE").status_code == 404
    assert cliente.post("/preguntar", data="pregunta=hola").status_code == 400  # solo JSON


# --- separar por cliente

def test_ia_solo_ve_al_cliente_elegido(dos_clientes):
    con = preguntas.conectar(OTRO_CLIENTE)
    assert con.execute("SELECT COUNT(*), MIN(receptor_rfc) FROM facturas").fetchone() == (2, OTRO_CLIENTE)
    assert con.execute("SELECT COUNT(*) FROM conceptos").fetchone() == (2,)
    # aunque el modelo pida la tabla original, el autorizador de SQLite no lo deja
    for sql in ("SELECT COUNT(*) FROM main.facturas", 'SELECT COUNT(*) FROM "main"."facturas"'):
        with pytest.raises(sqlite3.DatabaseError):
            con.execute(sql)
    con.close()
    with pytest.raises(ValueError):
        preguntas.conectar("X' OR 1=1 --")  # el RFC se valida antes de armar las vistas


def test_pdf_de_otro_cliente(dos_clientes, monkeypatch):
    modelo = ModeloFalso([("enviar_pdf", {"uuids": [FACTURA_01, FACTURA_OTRO]})])
    monkeypatch.setattr(preguntas.openai, "OpenAI", lambda: modelo)
    _, _, adjuntos = preguntas.responder("PDF de las dos", cliente=OTRO_CLIENTE)
    assert [a["uuid"] for a in adjuntos] == [FACTURA_OTRO]  # la de CLIENTE DEMO no se entrega
    # y el modelo recibió en sus instrucciones con qué cliente trabaja
    assert "OTRO CLIENTE DEMO" in modelo.recibido[0][0]["content"]


def test_filtro_por_cliente_en_la_web(cliente):
    datos = {"xml": [(p.open("rb"), p.name) for p in sorted(EJEMPLOS.glob("*.xml"))]}
    r = cliente.post("/subir", data=datos, headers={"X-Bot": "1"}).get_json()
    assert [c["rfc"] for c in r["clientes"]] == ["CDE150101AB1", OTRO_CLIENTE]

    r = cliente.post("/reto/resumen", json={"cliente": OTRO_CLIENTE}).get_json()
    assert "OTRO CLIENTE DEMO" in r["texto"] and len(r["consultas"][0]["filas"]) == 2
    r = cliente.post("/reto/clientes", json={}).get_json()
    assert [f[0] for f in r["consultas"][0]["filas"]] == ["CDE150101AB1", OTRO_CLIENTE]
    assert "2 facturas" in cliente.get(f"/estado?cliente={OTRO_CLIENTE}").get_json()["cifras"]

    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(cliente.get(f"/reporte_cfdi.xlsx?cliente={OTRO_CLIENTE}").data))
    assert wb["Detalle"].max_row == 3  # encabezado + 2 facturas
    assert cliente.post("/reto/detalle", json={"cliente": "no es rfc"}).status_code == 400


# --- subir XML desde el chat

def test_subir_xml(cliente):
    archivos = [(p.open("rb"), p.name) for p in sorted(EJEMPLOS.glob("*.xml"))] + [(BytesIO(b"hola"), "notas.txt")]
    r = cliente.post("/subir", data={"xml": archivos}, headers={"X-Bot": "1"}).get_json()
    assert "**3 archivos XML**" in r["texto"] and "**2 facturas cargadas**" in r["texto"]
    assert "notas.txt" in r["texto"]  # avisa que lo ignoró
    assert [(f[0], f[2]) for f in r["consultas"][0]["filas"]] == [("reenvio_factura_05.xml", "Duplicada")]
    assert len(list((config.SALIDA / "subidos").iterdir())) == 1  # cada subida en su propia carpeta


def test_subir_rechaza_peticiones_ajenas(cliente):
    archivo = {"xml": [(BytesIO(b"<x/>"), "a.xml")]}
    assert cliente.post("/subir", data=archivo).status_code == 400  # sin la cabecera X-Bot
    assert cliente.post("/subir", data={}, headers={"X-Bot": "1"}).status_code == 400  # sin archivos


def test_subir_nombre_peligroso(cliente):
    xml = (EJEMPLOS / "otro_cliente_01.xml").read_bytes()
    cliente.post("/subir", data={"xml": [(BytesIO(xml), "../../../fuera.xml")]}, headers={"X-Bot": "1"})
    (carpeta,) = (config.SALIDA / "subidos").iterdir()
    assert [p.name for p in carpeta.iterdir()] == ["fuera.xml"]  # se guardó dentro de su carpeta
    assert not (config.SALIDA.parent / "fuera.xml").exists()


def test_busqueda_sin_acentos(salida):
    # en la base el nombre viene sin acentos; la IA busca con norm() y lo encuentra igual
    con = preguntas.conectar()
    sql = "SELECT COUNT(*) FROM facturas WHERE norm(emisor_nombre) LIKE '%' || norm('Logísticos Tapatíos') || '%'"
    assert con.execute(sql).fetchone() == (3,)
    con.close()

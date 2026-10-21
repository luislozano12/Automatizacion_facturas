"""Pruebas de lo que pide el reto (reto.py y la carpeta revision/), con los XML de cfdis/."""
from pathlib import Path

import pytest
from openpyxl import load_workbook

import reto
from revision import config, reportes
from revision.base import cargar
from revision.lectura import Invalido, leer_cfdi

OTRO_CLIENTE = "OCD200101XY9"  # receptor de los XML de ejemplos/otro_cliente/


def test_extrae_los_campos():
    factura, conceptos, avisos = leer_cfdi(config.CFDIS / "factura_01.xml")
    assert factura["uuid"] == "A5AEC797-8306-D03B-F38B-2FFC80A4DF5A"
    assert (factura["emisor_rfc"], factura["emisor_nombre"]) == ("PAP180312KL4", "PAPELERIA DEL BAJIO DEMO")
    assert factura["fecha"] == "2026-09-28T09:56:00"
    assert (factura["subtotal"], factura["iva"], factura["total"]) == (6466.50, 1034.64, 7501.14)
    assert len(conceptos) == 1 and avisos == []


@pytest.mark.parametrize("archivo, tipo", [("factura_19.xml", "XML dañado"), ("factura_20.xml", "No es CFDI")])
def test_archivos_invalidos(archivo, tipo):
    with pytest.raises(Invalido) as e:
        leer_cfdi(config.CFDIS / archivo)
    assert e.value.tipo == tipo


def test_xml_malicioso(tmp_path):
    # "billion laughs": entidades que se expanden sin fin para agotar la memoria
    bomba = tmp_path / "bomba.xml"
    bomba.write_text('<?xml version="1.0"?><!DOCTYPE r [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;&a;">]><r>&b;</r>')
    with pytest.raises(Invalido) as e:
        leer_cfdi(bomba)
    assert e.value.tipo == "XML no permitido"


def test_carga_sin_detenerse(salida, consulta):
    assert consulta("SELECT COUNT(*) FROM facturas") == [(14,)]
    errores = {(Path(a).name, t) for a, t in consulta("SELECT archivo, tipo FROM incidencias WHERE nivel='ERROR'")}
    assert errores == {("factura_04.xml", "Duplicada"), ("factura_11.xml", "Duplicada"),
                       ("factura_19.xml", "XML dañado"), ("factura_20.xml", "No es CFDI")}
    log = (salida / "incidencias.log").read_text(encoding="utf-8")
    assert "factura_19.xml | XML dañado" in log and "Duplicada" in log


def test_cargar_dos_veces_da_lo_mismo(consulta):
    antes = consulta("SELECT * FROM facturas ORDER BY uuid"), consulta("SELECT archivo, tipo FROM incidencias ORDER BY 1, 2")
    cargar(config.CFDIS)
    despues = consulta("SELECT * FROM facturas ORDER BY uuid"), consulta("SELECT archivo, tipo FROM incidencias ORDER BY 1, 2")
    assert antes == despues


def test_resumen_de_mayor_a_menor(salida):
    _, filas = reportes.resumen()
    assert [f[0] for f in filas][:2] == ["PAP180312KL4", "ALI210622RT5"]
    assert [f[5] for f in filas] == sorted((f[5] for f in filas), reverse=True)
    pap = filas[0]
    assert pap[2] == 4 and pap[5] == 107825.92  # el duplicado (factura_04) no suma


def test_reportes_por_cliente(dos_clientes, consulta):
    assert consulta("SELECT COUNT(*) FROM facturas") == [(16,)]  # 14 + 2; el reenvío de factura_05 es duplicado
    _, clientes = reportes.resumen_clientes()
    assert [(c[0], c[2]) for c in clientes] == [("CDE150101AB1", 14), (OTRO_CLIENTE, 2)]
    _, filas = reportes.detalle(OTRO_CLIENTE)
    assert len(filas) == 2 and {f[7] for f in filas} == {OTRO_CLIENTE}
    _, filas = reportes.resumen(OTRO_CLIENTE)
    assert [f[0] for f in filas] == ["CON170405PL9", "SLT200915QW2"]
    _, filas = reportes.incidencias(carpeta=config.RAIZ / "ejemplos" / "otro_cliente")
    assert [(f[0], f[2]) for f in filas] == [("reenvio_factura_05.xml", "Duplicada")]


def test_reto_genera_excel_y_log(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "SALIDA", tmp_path)
    monkeypatch.setattr(config, "DB", tmp_path / "facturas.db")
    reto.main([])
    wb = load_workbook(tmp_path / "reporte_cfdi.xlsx")
    assert wb.sheetnames == ["Detalle", "Resumen por proveedor"]
    assert wb["Detalle"].max_row == 19  # título (3 filas) + encabezado + 14 facturas + total
    assert wb["Detalle"]["A4"].value == "UUID" and wb["Detalle"]["A19"].value == "TOTAL"
    assert wb["Resumen por proveedor"]["A5"].value == "PAP180312KL4"
    assert wb["Resumen por proveedor"]["F11"].value == "=SUM(F5:F10)"  # el total es fórmula
    assert "Duplicada" in (tmp_path / "incidencias.log").read_text(encoding="utf-8")


def test_reto_sin_carpeta(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "SALIDA", tmp_path)
    monkeypatch.setattr(config, "DB", tmp_path / "facturas.db")
    with pytest.raises(SystemExit, match="No existe la carpeta"):
        reto.main([str(tmp_path / "no-existe")])


def test_excel_del_reto_solo_con_su_carpeta(dos_clientes):
    # la base tiene además las facturas de otro cliente (subidas desde el bot), pero el Excel del reto no
    wb = load_workbook(reportes.excel(dos_clientes / "r.xlsx", carpeta=config.CFDIS))
    assert wb["Detalle"].max_row == 19  # título (3 filas) + encabezado + 14 facturas + total
    assert "Carpeta: cfdis" in wb["Detalle"]["A2"].value
